"""PPTX executive summary builder."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt


def _add_title(slide, title: str, subtitle: str | None = None) -> None:
    slide.shapes.title.text = title
    if subtitle and len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle


def _add_bullets(slide, bullets: list[str]) -> None:
    text_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.8), Inches(4.7))
    frame = text_box.text_frame
    frame.clear()
    for idx, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.size = Pt(22)


def _add_table(slide, df: pd.DataFrame, columns: list[str]) -> None:
    rows = min(len(df), 6) + 1
    table = slide.shapes.add_table(rows, len(columns), Inches(0.5), Inches(1.5), Inches(9.0), Inches(4.6)).table
    for idx, column in enumerate(columns):
        table.cell(0, idx).text = column
    for row_idx, (_, row) in enumerate(df[columns].head(6).iterrows(), start=1):
        for col_idx, column in enumerate(columns):
            table.cell(row_idx, col_idx).text = str(row[column])


def build_executive_summary_pptx(
    quality_df: pd.DataFrame,
    issues_df: pd.DataFrame,
    output_path: str | Path = "exports/executive_summary.pptx",
) -> Path:
    """Create a concise executive summary PowerPoint deck."""
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    _add_title(
        slide,
        "RetailOps Governance Intelligence",
        "Executive summary of data health, risk, and remediation priorities",
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    _add_title(slide, "Data Quality Scorecard")
    _add_table(
        slide,
        quality_df,
        ["table", "completeness", "consistency", "validity", "uniqueness", "composite_score"],
    )

    critical_high = 0 if issues_df.empty else int(issues_df["severity"].isin(["CRITICAL", "HIGH"]).sum())
    avg_quality = round(quality_df["composite_score"].mean(), 1)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    _add_title(slide, "Governance Snapshot")
    _add_bullets(
        slide,
        [
            f"Average governed table quality score: {avg_quality}/100.",
            f"Open governance issues generated: {len(issues_df)}.",
            f"Critical or high priority issues: {critical_high}.",
            "Primary remediation focus: improve completeness and resolve owned SLA items.",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    _add_title(slide, "Recommended Actions")
    _add_bullets(
        slide,
        [
            "Operationalize quality scoring as a pre-dashboard control.",
            "Assign data stewards to high-severity quality and PII issues.",
            "Publish the metadata catalog as the shared source of truth.",
            "Review stockout, churn, and anomaly signals in weekly business forums.",
        ],
    )

    prs.save(path)
    return path
