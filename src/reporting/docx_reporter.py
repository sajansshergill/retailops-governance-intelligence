"""DOCX report export module."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document


def _add_dataframe_table(document: Document, df: pd.DataFrame, columns: list[str], max_rows: int = 12) -> None:
    """Append a compact DataFrame preview table to a Word document."""
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Light Shading Accent 1"
    for idx, column in enumerate(columns):
        table.rows[0].cells[idx].text = column

    for _, row in df[columns].head(max_rows).iterrows():
        cells = table.add_row().cells
        for idx, column in enumerate(columns):
            cells[idx].text = str(row[column])


def build_business_requirements_docx(
    catalog_df: pd.DataFrame,
    quality_df: pd.DataFrame,
    issues_df: pd.DataFrame,
    output_path: str | Path = "exports/business_requirements.docx",
) -> Path:
    """Create a business requirements and governance summary document."""
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    document.add_heading("RetailOps Data Governance Requirements", level=0)
    document.add_paragraph(
        "This document summarizes governed retail data interfaces, quality "
        "expectations, ownership, PII handling, and open governance issues."
    )

    document.add_heading("Data Interfaces", level=1)
    interface_cols = ["asset_id", "data_type", "pii_type", "classification", "owner"]
    _add_dataframe_table(document, catalog_df, interface_cols, max_rows=18)

    document.add_heading("Quality Thresholds", level=1)
    quality_cols = ["table", "completeness", "consistency", "validity", "uniqueness", "composite_score"]
    _add_dataframe_table(document, quality_df, quality_cols, max_rows=10)

    document.add_heading("Governance Issue Summary", level=1)
    if issues_df.empty:
        document.add_paragraph("No open governance issues were generated for this sample run.")
    else:
        issue_cols = ["source_table", "rule_name", "severity", "owner", "sla_due"]
        _add_dataframe_table(document, issues_df, issue_cols, max_rows=15)

    document.add_heading("Acceptance Criteria", level=1)
    for criterion in [
        "Core source tables are profiled before BI consumption.",
        "PII columns are tagged and classified in the catalog.",
        "Tables below quality thresholds produce owned governance issues.",
        "Lineage is available for dashboard and executive reporting outputs.",
    ]:
        document.add_paragraph(criterion, style="List Bullet")

    document.save(path)
    return path
